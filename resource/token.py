# -*- coding: utf-8 -*-
"""
自动化脚本：从两个目录自动查找文档 → 匹配告警指纹（仅保留挂起原因非空）
依赖：pip install pandas openpyxl
"""

import pandas as pd
import os
import glob

# ==================== 配置区 ====================

# 两个文档所在的目录（改成你的实际路径）
A_DIR = r"/Users/hjy/监控交付/script/token/告警"          # A文档所在目录
B_DIR = r"/Users/hjy/监控交付/script/token/挂起"          # B文档所在目录

# 文件名关键字（可选，用于在目录中精准定位文件；留空则取目录下第一个xlsx）
A_KEYWORD = ""                    # 例如 "告警" 或 "alarm"
B_KEYWORD = ""                    # 例如 "设备" 或 "device"

# A文档列名
A_COL_NAME   = "告警名称"
A_COL_DEVICE = "告警设备"
A_COL_FINGER = "告警指纹"

# B文档列名
B_COL_TYPE   = "告警类型"
B_COL_DEVICE = "设备编号"
B_COL_REASON = "挂起原因"

SEP = "|"                         # 拼接分隔符

# 输出目录和文件名
OUTPUT_DIR  = r"/Users/hjy/监控交付/script/token/指纹"      # 输出目录
OUTPUT_FILE = "匹配结果.xlsx"      # 输出文件名

# ==================== 新增配置：每列最大行数 ====================
MAX_ROWS_PER_COLUMN = 1000         # 每列最多行数
# ================================================================

# ==================== 以下逻辑无需修改 ====================

def find_excel(dir_path, keyword=""):
    """在目录中查找xlsx文件，可按关键字过滤"""
    if not os.path.isdir(dir_path):
        raise FileNotFoundError(f"目录不存在: {dir_path}")

    # 查找所有xlsx文件（不包含 ~$ 开头的临时文件）
    files = glob.glob(os.path.join(dir_path, "*.xlsx"))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]

    if not files:
        raise FileNotFoundError(f"目录中未找到xlsx文件: {dir_path}")

    # 按关键字过滤
    if keyword:
        matched = [f for f in files if keyword in os.path.basename(f)]
        if not matched:
            print(f"警告: 目录 {dir_path} 中未找到包含 '{keyword}' 的文件，将使用全部文件")
        else:
            files = matched

    # 返回找到的第一个文件（如多个，提示）
    if len(files) > 1:
        print(f"提示: 目录 {dir_path} 中找到多个文件，默认使用第一个:")
        for f in files:
            print(f"    - {os.path.basename(f)}")

    return files[0]


def load_sheet(path, sheet=""):
    """读取Excel，sheet为空则取第一个sheet"""
    if sheet:
        return pd.read_excel(path, sheet_name=sheet, dtype=str)
    return pd.read_excel(path, sheet_name=0, dtype=str)


def write_columnar_excel(df_output, output_path, max_rows=1000):
    """
    将数据按列写入Excel，每列最多max_rows行，
    超出部分隔一列后从新的一列继续填充。
    
    数据结构：
    - 每组列包含4列：告警名称、告警设备、告警指纹、挂起原因
    - 组与组之间隔1个空列
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    ws = wb.active
    ws.title = "匹配结果"
    
    headers = ["告警名称", "告警设备", "告警指纹", "挂起原因"]
    total_rows = len(df_output)
    
    # 计算需要多少组列
    num_groups = (total_rows + max_rows - 1) // max_rows  # 向上取整
    print(f"总记录数: {total_rows}，每列最多 {max_rows} 行，需要 {num_groups} 组列")
    
    for group_idx in range(num_groups):
        start_row = group_idx * max_rows
        end_row = min(start_row + max_rows, total_rows)
        
        # 当前组的起始列位置（每组4列 + 1个空列间隔）
        start_col = group_idx * 5 + 1  # 组间隔：4列数据 + 1空列
        
        # 写入表头
        for col_idx, header in enumerate(headers):
            ws.cell(row=1, column=start_col + col_idx, value=header)
        
        # 写入数据
        for row_offset, data_idx in enumerate(range(start_row, end_row)):
            excel_row = row_offset + 2  # 从第2行开始（第1行是表头）
            ws.cell(row=excel_row, column=start_col, value=df_output.iloc[data_idx]["告警名称"])
            ws.cell(row=excel_row, column=start_col + 1, value=df_output.iloc[data_idx]["告警设备"])
            ws.cell(row=excel_row, column=start_col + 2, value=df_output.iloc[data_idx]["告警指纹"])
            ws.cell(row=excel_row, column=start_col + 3, value=df_output.iloc[data_idx]["挂起原因"])
    
    # 调整列宽（可选）
    for col_idx in range(1, num_groups * 5 + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 30
    
    wb.save(output_path)
    print(f"结果已保存到: {output_path}")


def main():
    # 自动查找文件
    a_path = find_excel(A_DIR, A_KEYWORD)
    b_path = find_excel(B_DIR, B_KEYWORD)

    print(f"A文档: {a_path}")
    print(f"B文档: {b_path}")
    print("-" * 50)

    df_a = load_sheet(a_path)
    df_b = load_sheet(b_path)

    print(f"A文档行数: {len(df_a)}")
    print(f"B文档行数: {len(df_b)}")

    # 拼接列
    df_a["P"] = df_a[A_COL_NAME].fillna("").str.strip() + SEP + df_a[A_COL_DEVICE].fillna("").str.strip()
    df_b["M"] = df_b[B_COL_TYPE].fillna("").str.strip() + SEP + df_b[B_COL_DEVICE].fillna("").str.strip()

    # B文档M去重
    df_b_unique = df_b.drop_duplicates(subset=["M"], keep="first")

    # 重命名B文档的挂起原因，避免与A文档同名列冲突
    df_b_unique = df_b_unique.rename(columns={B_COL_REASON: "挂起原因_B"})

    # 左连接匹配
    df_result = df_a.merge(
        df_b_unique[["M", "挂起原因_B"]],
        left_on="P",
        right_on="M",
        how="left"
    )

    # 挂起原因非空过滤
    df_result["挂起原因_B"] = df_result["挂起原因_B"].fillna("").str.strip()
    df_filtered = df_result[df_result["挂起原因_B"] != ""].copy()

    print(f"匹配到挂起原因非空的记录数: {len(df_filtered)}")

    # 输出
    df_output = df_filtered[[A_COL_NAME, A_COL_DEVICE, A_COL_FINGER, "挂起原因_B"]]
    df_output = df_output.rename(columns={"挂起原因_B": "挂起原因"})

    # 确保输出目录存在
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, OUTPUT_FILE)

    # 使用新的按列写入方式
    write_columnar_excel(df_output, output_path, MAX_ROWS_PER_COLUMN)

    fingers = df_output[A_COL_FINGER].dropna().tolist()
    print(f"\n共 {len(fingers)} 条告警指纹：")
    for f in fingers:
        print(f"  {f}")


if __name__ == "__main__":
    main()
