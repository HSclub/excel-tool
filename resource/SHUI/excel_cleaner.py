# -*- coding: utf-8 -*-

"""
Excel 设备编号自动填充工具

功能：

1. 从 input 目录读取设备清单。
2. 从 output 目录读取模板。
3. 复制模板 Sheet。
4. 按 groups.match 筛选设备。
5. template_find 定位设备类型标题。
6. auto_position 支持设备块连续向下填充。
7. 当前设备数量超过模板预留空间时：
   在下一设备类型标题之前自动插入行。
8. 自动复制模板行格式。
9. 不覆盖后续设备类型标题。
10. 支持 ${project_name} 配置变量。
11. 每个 output_sheet 单独保存为 Excel。
"""

import copy
import os
import re
import sys
from pathlib import Path

import pandas as pd
import yaml

from openpyxl import load_workbook
from openpyxl.utils import (
    column_index_from_string,
)
from openpyxl.styles import PatternFill

# ============================================================
# 基础工具
# ============================================================

def normalize(value):
    """Excel 单元格值统一转换为可比较字符串。"""

    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    return str(value).strip()


# ============================================================
# 匹配
# ============================================================

def value_match(actual, expected):
    """
    支持：

    expected: "xxx"

    expected:
        contains: "xxx"

    expected:
        startswith: "xxx"

    expected:
        endswith: "xxx"

    expected:
        regex: "xxx"

    expected:
        in:
          - "xxx"
          - "yyy"
    """

    actual = normalize(actual)

    if isinstance(expected, dict):

        if "contains" in expected:
            return normalize(expected["contains"]) in actual

        if "startswith" in expected:
            return actual.startswith(
                normalize(expected["startswith"])
            )

        if "endswith" in expected:
            return actual.endswith(
                normalize(expected["endswith"])
            )

        if "regex" in expected:
            try:
                return re.search(
                    str(expected["regex"]),
                    actual
                ) is not None
            except re.error as exc:
                raise ValueError(
                    f"正则表达式错误：{expected['regex']}"
                ) from exc

        if "in" in expected:
            return actual in {
                normalize(x)
                for x in expected["in"]
            }

        raise ValueError(
            f"不支持的匹配方式：{expected}"
        )

    if isinstance(expected, list):

        return actual in {
            normalize(x)
            for x in expected
        }

    return actual == normalize(expected)


def row_matches(row, conditions):
    """判断一条设备记录是否满足 match。"""

    for field, expected in (conditions or {}).items():

        if field not in row.index:
            return False

        if not value_match(
            row[field],
            expected
        ):
            return False

    return True


# ============================================================
# Excel 样式
# ============================================================

def copy_cell_style(src, dst):
    """复制单元格格式。"""

    if src.has_style:
        dst._style = copy.copy(src._style)

    if src.number_format:
        dst.number_format = src.number_format

    if src.alignment:
        dst.alignment = copy.copy(src.alignment)

    if src.protection:
        dst.protection = copy.copy(src.protection)

    if src.font:
        dst.font = copy.copy(src.font)

    if src.fill:
        dst.fill = copy.copy(src.fill)

    if src.border:
        dst.border = copy.copy(src.border)


def copy_row_style(ws, source_row, target_row):
    """
    把 source_row 的格式复制到 target_row。

    不复制内容。
    """

    if source_row in ws.row_dimensions:

        src_dim = ws.row_dimensions[source_row]
        dst_dim = ws.row_dimensions[target_row]

        dst_dim.height = src_dim.height
        dst_dim.hidden = src_dim.hidden
        dst_dim.outlineLevel = src_dim.outlineLevel
        dst_dim.collapsed = src_dim.collapsed

    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)

        copy_cell_style(src, dst)

        # 强制新插入行背景为白色
        dst.fill = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF"
        )


def clear_row_values(ws, row):
    """清除整行内容，但保留格式。"""

    for col in range(1, ws.max_column + 1):
        ws.cell(row, col).value = None


# ============================================================
# 查找模板标题
# ============================================================

def find_cell(ws, column, value, start_row=1):
    """
    在指定列查找指定值。

    start_row 可以控制搜索起始位置。
    """

    col_idx = column_index_from_string(column)

    for row in range(
        start_row,
        ws.max_row + 1
    ):

        cell_value = normalize(
            ws.cell(row, col_idx).value
        )

        if cell_value == normalize(value):
            return row

    return None


def find_template_anchor(ws, template_find):
    """根据 template_find 查找模板标题。"""

    if not template_find:
        return None

    column = template_find["column"]
    value = template_find["value"]

    return find_cell(
        ws,
        column,
        value
    )


# ============================================================
# 输入文件
# ============================================================

def safe_filename(name):
    """处理 Windows 文件名非法字符。"""

    return re.sub(
        r'[\\/:*?"<>|]',
        "_",
        str(name)
    ).strip()


def resolve_template(output_dir, template_file):
    """模板必须从 output_dir 获取。"""

    path = Path(output_dir) / template_file

    if path.exists():
        return path

    return None


def select_input_files(
    input_dir,
    input_file=None,
    input_file_contains=None,
    input_file_regex=None
):
    """
    选择输入 Excel。

    优先级：

    1. input_file
    2. input_file_contains
    3. input_file_regex
    4. 全部 Excel
    """

    input_path = Path(input_dir)

    if not input_path.exists():
        raise FileNotFoundError(
            f"输入目录不存在：{input_path}"
        )

    files = sorted(
        p
        for p in input_path.iterdir()
        if (
            p.is_file()
            and p.suffix.lower()
            in {".xlsx", ".xlsm"}
            and not p.name.startswith("~$")
        )
    )

    if not files:
        raise FileNotFoundError(
            f"输入目录没有找到 Excel：{input_path}"
        )

    # --------------------------------------------------------
    # 1. 精确文件名
    # --------------------------------------------------------

    if input_file:

        target = input_path / str(input_file)

        if not target.exists():
            available = "\n".join(
                f"  - {p.name}"
                for p in files
            )

            raise FileNotFoundError(
                f"指定的输入文件不存在：{target}\n"
                f"当前 input 目录中的 Excel：\n"
                f"{available}"
            )

        return [target]

    # --------------------------------------------------------
    # 2. 模糊包含
    # --------------------------------------------------------

    if input_file_contains:

        keyword = str(
            input_file_contains
        ).strip()

        matched = [
            p
            for p in files
            if keyword in p.name
        ]

        if not matched:

            available = "\n".join(
                f"  - {p.name}"
                for p in files
            )

            raise FileNotFoundError(
                f"没有找到文件名包含【{keyword}】的 Excel。\n"
                f"当前 input 目录中的 Excel：\n"
                f"{available}"
            )

        return matched

    # --------------------------------------------------------
    # 3. 正则
    # --------------------------------------------------------

    if input_file_regex:

        try:
            pattern = re.compile(
                str(input_file_regex)
            )
        except re.error as exc:

            raise ValueError(
                f"input_file_regex 正则表达式错误："
                f"{input_file_regex}"
            ) from exc

        matched = [
            p
            for p in files
            if pattern.search(p.name)
        ]

        if not matched:

            available = "\n".join(
                f"  - {p.name}"
                for p in files
            )

            raise FileNotFoundError(
                f"没有找到符合正则【{input_file_regex}】的 Excel。\n"
                f"当前 input 目录中的 Excel：\n"
                f"{available}"
            )

        return matched

    # --------------------------------------------------------
    # 4. 没有限制
    # --------------------------------------------------------

    return files


def read_input_files(
    input_dir,
    input_sheet,
    input_file=None,
    input_file_contains=None,
    input_file_regex=None
):
    """读取输入 Excel。"""

    files = select_input_files(
        input_dir=input_dir,
        input_file=input_file,
        input_file_contains=input_file_contains,
        input_file_regex=input_file_regex
    )

    print(
        f"选中的输入 Excel：{len(files)} 个"
    )

    for file in files:
        print(f"  - {file.name}")

    all_frames = []

    for file in files:

        print(
            f"读取输入文件：{file}"
        )

        try:

            df = pd.read_excel(
                file,
                sheet_name=input_sheet,
                dtype=str
            )

        except ValueError as exc:

            raise ValueError(
                f"文件【{file.name}】找不到 Sheet："
                f"{input_sheet}"
            ) from exc

        df = df.fillna("")

        df.columns = [
            normalize(c)
            for c in df.columns
        ]

        all_frames.append(df)

    return pd.concat(
        all_frames,
        ignore_index=True
    )


# ============================================================
# 字段检查
# ============================================================

def validate_required_columns(df, config):

    required = set(
        config.get(
            "require_non_empty",
            []
        )
    )

    for group in config.get(
        "groups",
        []
    ):

        value_column = group.get(
            "value_column"
        )

        if value_column:
            required.add(value_column)

        for field in (
            group.get("match") or {}
        ).keys():

            required.add(field)

    missing = [
        x
        for x in required
        if x not in df.columns
    ]

    if missing:

        raise ValueError(
            "输入 Excel 缺少字段："
            + "、".join(missing)
            + f"\n实际字段：{list(df.columns)}"
        )


# ============================================================
# 筛选设备
# ============================================================

def filter_devices(
    df,
    group,
    require_non_empty
):
    """按照 group.match 筛选设备。"""

    result = df.copy()

    # --------------------------------------------------------
    # 必填字段
    # --------------------------------------------------------

    for field in require_non_empty:

        if field in result.columns:

            result = result[
                result[field].map(normalize)
                != ""
            ]

    # --------------------------------------------------------
    # match
    # --------------------------------------------------------

    match = group.get("match") or {}

    if match:

        mask = result.apply(
            lambda row:
                row_matches(
                    row,
                    match
                ),
            axis=1
        )

        result = result[mask]

    return result


def get_values(df, value_column):
    """取得设备完整编号。"""

    values = []

    for value in df[value_column].tolist():

        value = normalize(value)

        if value:
            values.append(value)

    return values


# ============================================================
# 查找下一个模板块
# ============================================================

def find_next_group_anchor(
    ws,
    current_group_index,
    groups,
    current_anchor_row
):
    """
    找当前 group 后面的下一个模板标题。

    例如：

        当前：
        输入断路器-市电

        下一个：
        输入断路器-逆变器

    只有后面的 group 配置了 template_find，
    才可以准确找到下一块模板。
    """

    for index in range(
        current_group_index + 1,
        len(groups)
    ):

        next_group = groups[index]

        template_find = next_group.get(
            "template_find"
        )

        if not template_find:
            continue

        row = find_cell(
            ws,
            template_find["column"],
            template_find["value"],
            start_row=current_anchor_row + 1
        )

        if row is not None:
            return row

    return None


# ============================================================
# 在指定位置插入行
# ============================================================

def insert_rows_with_style(
    ws,
    row,
    count,
    style_row
):
    """
    在 row 前插入 count 行。

    新行复制 style_row 格式。
    """

    if count <= 0:
        return

    print(
        f"    自动插入 {count} 行，"
        f"位置：第 {row} 行"
    )

    ws.insert_rows(
        row,
        amount=count
    )

    for r in range(
        row,
        row + count
    ):

        copy_row_style(
            ws,
            style_row,
            r
        )

        clear_row_values(
            ws,
            r
        )


# ============================================================
# 写入纵向数据
# ============================================================

def write_vertical_dynamic(
    ws,
    start_row,
    column,
    values,
    style_row,
    next_anchor_row=None,
    blank_rows_after=0
):
    """
    向下写设备编号。

    如果存在下一块模板：

        当前数据区域
        ↓
        下一块模板

    数据不够时，直接在下一块模板之前插入行。

    这样不会覆盖下一块模板。
    """

    if not values:
        return start_row - 1

    col_idx = column_index_from_string(
        column
    )

    count = len(values)

    # ========================================================
    # 情况一：
    # 后面有下一块模板
    # ========================================================

    if next_anchor_row is not None:

        # 当前数据最多可以使用到下一标题前
        available_rows = (
            next_anchor_row
            - start_row
            - blank_rows_after
        )

        if available_rows < 0:
            available_rows = 0

        # 需要增加多少行
        insert_count = max(
            0,
            count - available_rows
        )

        if insert_count > 0:

            insert_rows_with_style(
                ws=ws,
                row=next_anchor_row,
                count=insert_count,
                style_row=style_row
            )

            # 下一模板标题已经被向下移动
            next_anchor_row += insert_count

    # ========================================================
    # 情况二：
    # 最后一块模板，没有下一标题
    # ========================================================

    else:

        last_required_row = (
            start_row + count - 1
        )

        if last_required_row > ws.max_row:

            insert_count = (
                last_required_row
                - ws.max_row
            )

            old_max_row = ws.max_row

            ws.insert_rows(
                ws.max_row + 1,
                amount=insert_count
            )

            for row in range(
                old_max_row + 1,
                last_required_row + 1
            ):

                copy_row_style(
                    ws,
                    style_row,
                    row
                )

                clear_row_values(
                    ws,
                    row
                )

    # ========================================================
    # 写入数据
    # ========================================================

    for index, value in enumerate(values):

        row = start_row + index

        ws.cell(
            row,
            col_idx
        ).value = value

    return (
        start_row
        + count
        - 1
    )


# ============================================================
# 处理一个 group
# ============================================================

def process_group(
    ws,
    group,
    group_index,
    groups,
    df,
    require_non_empty,
    state
):
    """
    处理一个设备 group。

    核心逻辑：

    1. 找当前设备类型标题。
    2. 找下一设备类型标题。
    3. 从当前标题下一行开始写。
    4. 如果数据超出预留空间：
       在下一设备标题之前插入行。
    5. 下一设备标题自动下移。
    """

    group = dict(group)

    values = get_values(
        filter_devices(
            df,
            group,
            require_non_empty
        ),
        group["value_column"]
    )

    name = group.get(
        "name",
        "未命名"
    )

    print()
    print(
        f"[{name}]"
    )

    print(
        f"筛选到设备：{len(values)} 个"
    )

    if values:

        for value in values:
            print(
                f"  {value}"
            )

    else:

        print(
            "  没有匹配设备，跳过写入。"
        )

        return state

    direction = group.get(
        "direction",
        "down"
    )

    if direction != "down":

        raise ValueError(
            f"当前版本主要用于纵向填充，"
            f"group【{name}】direction={direction}"
        )

    # ========================================================
    # 找当前模板标题
    # ========================================================

    template_find = group.get(
        "template_find"
    )

    current_anchor_row = None

    if template_find:

        current_anchor_row = find_template_anchor(
            ws,
            template_find
        )

    # ========================================================
    # auto_position
    #
    # 如果是 auto_position：
    # 数据从上一组数据结束位置继续。
    #
    # 但是我们仍然需要找到当前模板标题，
    # 用于确定扩容位置和样式。
    # ========================================================

    if group.get("auto_position"):

        if state.get("last_row") is None:

            raise ValueError(
                f"group【{name}】设置了 "
                f"auto_position: true，"
                f"但前面没有数据块。"
            )

        start_row = (
            state["last_row"] + 1
        )

        # 当前 group 的标题如果存在，
        # 那么真正的数据起始位置应该是标题下一行。
        if current_anchor_row is not None:

            expected_start = (
                current_anchor_row + 1
            )

            # 如果上一组数据因为扩容，
            # 已经正好延伸到当前标题前，
            # 使用上一组结束位置。
            #
            # 如果中间有模板空白区域，
            # 则优先使用当前标题下一行。
            if expected_start > start_row:
                start_row = expected_start

    else:

        if current_anchor_row is None:

            raise ValueError(
                f"Sheet【{ws.title}】中找不到模板字段："
                f"{template_find}"
            )

        start_row = (
            current_anchor_row + 1
        )

    # ========================================================
    # 如果找不到当前标题
    # ========================================================

    if current_anchor_row is None:

        raise ValueError(
            f"group【{name}】必须配置有效的 "
            f"template_find。"
        )

    # ========================================================
    # 当前模板行作为新增数据行样式
    # ========================================================

    style_row = current_anchor_row

    # ========================================================
    # 查找下一块模板标题
    # ========================================================

    next_anchor_row = find_next_group_anchor(
        ws=ws,
        current_group_index=group_index,
        groups=groups,
        current_anchor_row=current_anchor_row
    )

    if next_anchor_row:

        print(
            f"    当前标题行：{current_anchor_row}"
        )

        print(
            f"    下一模板标题行：{next_anchor_row}"
        )

    else:

        print(
            f"    当前标题行：{current_anchor_row}"
        )

        print(
            "    当前是最后一个模板块"
        )

    # ========================================================
    # 预留空白行
    # ========================================================

    blank_rows_after = int(
        group.get(
            "blank_rows_after",
            0
        ) or 0
    )

    # ========================================================
    # 写入
    # ========================================================

    last_row = write_vertical_dynamic(
        ws=ws,
        start_row=start_row,
        column=group["output_column"],
        values=values,
        style_row=style_row,
        next_anchor_row=next_anchor_row,
        blank_rows_after=blank_rows_after
    )

    # ========================================================
    # 保存状态
    # ========================================================

    state["last_row"] = last_row

    state["style_row"] = style_row

    # blank_rows_after 真正占用位置
    if blank_rows_after > 0:

        blank_start = (
            last_row + 1
        )

        blank_end = (
            last_row
            + blank_rows_after
        )

        # 如果空白区域已经超出 Sheet，
        # 增加行
        if blank_end > ws.max_row:

            insert_count = (
                blank_end
                - ws.max_row
            )

            old_max_row = ws.max_row

            ws.insert_rows(
                ws.max_row + 1,
                amount=insert_count
            )

            for row in range(
                old_max_row + 1,
                blank_end + 1
            ):

                copy_row_style(
                    ws,
                    style_row,
                    row
                )

        for row in range(
            blank_start,
            blank_end + 1
        ):

            clear_row_values(
                ws,
                row
            )

        state["last_row"] = blank_end

    return state


# ============================================================
# YAML 变量
# ============================================================

def expand_config_variables(
    value,
    variables
):
    """
    支持：

    project_name: "嘉兴洪福模组12"

    input_file_contains: "${project_name}"

    output_sheet: "${project_name}"
    """

    if isinstance(value, str):

        pattern = re.compile(
            r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}"
        )

        def replace(match):

            key = match.group(1)

            if key not in variables:

                raise ValueError(
                    f"配置引用了未定义的变量："
                    f"${{{key}}}"
                )

            return str(
                variables[key]
            )

        return pattern.sub(
            replace,
            value
        )

    if isinstance(value, list):

        return [
            expand_config_variables(
                item,
                variables
            )
            for item in value
        ]

    if isinstance(value, dict):

        return {
            key: expand_config_variables(
                val,
                variables
            )
            for key, val in value.items()
        }

    return value


# ============================================================
# 创建输出 Excel
# ============================================================

def create_output_for_group(
    template_path,
    template_sheet,
    output_sheet,
    df,
    config,
    groups,
    output_dir
):
    """
    创建一个独立的输出 Excel。
    """

    keep_vba = (
        template_path.suffix.lower()
        == ".xlsm"
    )

    wb = load_workbook(
        template_path,
        keep_vba=keep_vba
    )

    # --------------------------------------------------------
    # 找模板 Sheet
    # --------------------------------------------------------

    if template_sheet:

        if template_sheet not in wb.sheetnames:

            raise ValueError(
                f"模板 Sheet 不存在："
                f"{template_sheet}\n"
                f"当前 Sheet："
                f"{wb.sheetnames}"
            )

        source_ws = wb[
            template_sheet
        ]

    else:

        source_ws = wb[
            wb.sheetnames[0]
        ]

    # --------------------------------------------------------
    # 删除同名输出 Sheet
    # --------------------------------------------------------

    if output_sheet in wb.sheetnames:

        del wb[
            output_sheet
        ]

    # --------------------------------------------------------
    # 复制模板 Sheet
    # --------------------------------------------------------

    ws = wb.copy_worksheet(
        source_ws
    )

    ws.title = output_sheet

    # --------------------------------------------------------
    # 状态
    # --------------------------------------------------------

    state = {
        "last_row": None,
        "style_row": None,
    }

    require_non_empty = config.get(
        "require_non_empty",
        []
    )

    # --------------------------------------------------------
    # 按 group 顺序处理
    # --------------------------------------------------------

    for index, group in enumerate(groups):

        process_group(
            ws=ws,
            group=group,
            group_index=index,
            groups=groups,
            df=df,
            require_non_empty=require_non_empty,
            state=state
        )

    # --------------------------------------------------------
    # 保存
    # --------------------------------------------------------

    output_dir = Path(
        output_dir
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    ext = template_path.suffix

    output_name = (
        safe_filename(output_sheet)
        + ext
    )

    output_path = (
        output_dir
        / output_name
    )

    wb.save(
        output_path
    )

    return output_path


# ============================================================
# main
# ============================================================

def main():

    print("=" * 70)
    print("Excel 设备编号自动填充工具")
    print("=" * 70)

    config_path = Path(
        "config.yaml"
    )

    if not config_path.exists():

        print(
            f"[错误] 找不到 config.yaml："
            f"{config_path.resolve()}"
        )

        raise Exception("处理失败")

    # --------------------------------------------------------
    # 读取配置
    # --------------------------------------------------------

    with config_path.open(
        "r",
        encoding="utf-8"
    ) as f:

        config = yaml.safe_load(
            f
        ) or {}

    # --------------------------------------------------------
    # YAML 变量
    # --------------------------------------------------------

    variables = {
        key: value
        for key, value in config.items()
        if (
            isinstance(key, str)
            and isinstance(
                value,
                (
                    str,
                    int,
                    float,
                    bool
                )
            )
        )
    }

    try:

        config = expand_config_variables(
            config,
            variables
        )

    except Exception as exc:

        print(
            f"[错误] 配置变量解析失败：{exc}"
        )

        raise Exception("处理失败")

    # --------------------------------------------------------
    # 基础配置
    # --------------------------------------------------------

    input_dir = config.get(
        "input_dir",
        "input"
    )

    output_dir = config.get(
        "output_dir",
        "output"
    )

    template_file = config.get(
        "template_file"
    )

    template_sheet = config.get(
        "template_sheet"
    )

    input_sheet = config.get(
        "input_sheet",
        "设备清单"
    )

    input_file = config.get(
        "input_file"
    )

    input_file_contains = config.get(
        "input_file_contains"
    )

    input_file_regex = config.get(
        "input_file_regex"
    )

    global_output_sheet = config.get(
        "output_sheet"
    )

    groups = config.get(
        "groups",
        []
    )

    # --------------------------------------------------------
    # 基础检查
    # --------------------------------------------------------

    if not template_file:

        print(
            "[错误] config.yaml 没有配置 "
            "template_file"
        )

        raise Exception("处理失败")

    if not groups:

        print(
            "[错误] config.yaml 没有配置 groups"
        )

        raise Exception("处理失败")

    # --------------------------------------------------------
    # 读取输入 Excel
    # --------------------------------------------------------

    try:

        df = read_input_files(
            input_dir=input_dir,
            input_sheet=input_sheet,
            input_file=input_file,
            input_file_contains=input_file_contains,
            input_file_regex=input_file_regex
        )

    except Exception as exc:

        print(
            f"[错误] 读取输入 Excel 失败：{exc}"
        )

        raise Exception("处理失败")

    print()
    print(
        f"设备清单总行数：{len(df)}"
    )

    print(
        f"输入字段：{list(df.columns)}"
    )

    # --------------------------------------------------------
    # 检查字段
    # --------------------------------------------------------

    try:

        validate_required_columns(
            df,
            config
        )

    except Exception as exc:

        print(
            f"[错误] 输入字段检查失败：{exc}"
        )

        raise Exception("处理失败")

    # --------------------------------------------------------
    # 模板
    # --------------------------------------------------------

    template_path = resolve_template(
        output_dir,
        template_file
    )

    if template_path is None:

        print(
            f"\n[错误] 找不到模板文件："
            f"{Path(output_dir) / template_file}"
        )

        print(
            "请把模板 Excel 放到 output 目录。"
        )

        raise Exception("处理失败")

    print()
    print(
        f"模板文件："
        f"{template_path.resolve()}"
    )

    # --------------------------------------------------------
    # 按 output_sheet 分组
    # --------------------------------------------------------

    grouped = {}

    for group in groups:

        sheet_name = (
            group.get("output_sheet")
            or global_output_sheet
        )

        if not sheet_name:

            print(
                f"[错误] group【"
                f"{group.get('name', '未命名')}"
                f"】没有配置 output_sheet，"
                f"且全局没有 output_sheet。"
            )

            raise Exception("处理失败")

        grouped.setdefault(
            sheet_name,
            []
        ).append(group)

    # --------------------------------------------------------
    # 生成
    # --------------------------------------------------------

    success = 0

    for sheet_name, sheet_groups in grouped.items():

        print()
        print("=" * 70)
        print(
            f"开始生成：{sheet_name}"
        )
        print("=" * 70)

        try:

            output_path = (
                create_output_for_group(
                    template_path=template_path,
                    template_sheet=template_sheet,
                    output_sheet=sheet_name,
                    df=df,
                    config=config,
                    groups=sheet_groups,
                    output_dir=output_dir
                )
            )

            print()
            print(
                f"[成功] 输出文件："
                f"{output_path.resolve()}"
            )

            success += 1

        except Exception as exc:

            print()
            print(
                f"[错误] 生成【{sheet_name}】失败："
                f"{exc}"
            )

    # --------------------------------------------------------
    # 完成
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print(
        f"处理完成：成功 {success} 个，"
        f"失败 {len(grouped) - success} 个"
    )
    print("=" * 70)


if __name__ == "__main__":
    main()