from pathlib import Path
import sys
import re

import pandas as pd
import yaml
from openpyxl import load_workbook


# ==========================================================
# 基础目录
# ==========================================================

BASE_DIR = Path(__file__).resolve().parent


# ==========================================================
# 工具函数
# ==========================================================

def clean_value(value):
    """
    统一处理 Excel 单元格值

    None / NaN -> ""
    其他 -> 字符串并去掉首尾空格
    """

    if pd.isna(value):
        return ""

    return str(value).strip()


# ==========================================================
# 配置变量解析
# ==========================================================

def expand_variables(value, variables):
    """
    解析配置中的 ${变量名}

    例如：

    project_name: "乌兰察布乌拉模组44"
    input_file_contains: "${project_name}"
    output_sheet: "${project_name}"
    output_file: "${project_name}.xlsx"

    会自动变成：

    input_file_contains = "乌兰察布乌拉模组44"
    output_sheet = "乌兰察布乌拉模组44"
    output_file = "乌兰察布乌拉模组44.xlsx"
    """

    if not isinstance(value, str):
        return value

    pattern = r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}"

    def replace(match):

        name = match.group(1)

        if name not in variables:
            print(
                f"[警告] 未定义配置变量：${{{name}}}"
            )

            return match.group(0)

        return str(variables[name])

    return re.sub(
        pattern,
        replace,
        value
    )


def expand_config_variables(config):
    """
    递归解析整个 config 中的 ${变量名}
    """

    # 顶层配置作为变量来源
    variables = {
        key: value
        for key, value in config.items()
        if isinstance(
            value,
            (
                str,
                int,
                float,
                bool
            )
        )
    }

    def recursive(value):

        # 字符串
        if isinstance(value, str):

            return expand_variables(
                value,
                variables
            )

        # 字典
        if isinstance(value, dict):

            return {
                key: recursive(val)
                for key, val in value.items()
            }

        # 列表
        if isinstance(value, list):

            return [
                recursive(item)
                for item in value
            ]

        # 其他类型
        return value

    return recursive(config)


# ==========================================================
# 判断一行是否符合筛选条件
# ==========================================================

def row_matches(row, conditions):
    """
    判断一行数据是否满足 match 条件。

    默认所有条件都是 AND。

    例如：

    match:
      应用类型: "输入断路器"
      所属整机类型: "一体冷源"

    表示：

    应用类型 = 输入断路器
    AND
    所属整机类型 = 一体冷源


    如果配置的是列表：

    应用类型:
      - "输入断路器"
      - "输出断路器"

    表示 OR。
    """

    for column, expected in conditions.items():

        actual = clean_value(
            row.get(
                column,
                ""
            )
        )

        # ==================================================
        # 如果配置的是列表
        # ==================================================

        if isinstance(expected, list):

            expected_values = [
                clean_value(value)
                for value in expected
            ]

            if actual not in expected_values:
                return False

        else:

            expected_value = clean_value(
                expected
            )

            if actual != expected_value:
                return False

    return True


# ==========================================================
# 查找输入 Excel
# ==========================================================

def get_input_excel(
    input_dir,
    input_file_contains=None
):
    """
    从 input 目录查找 Excel。

    如果配置：

        input_file_contains: "乌兰察布乌拉模组44"

    则查找：

        xxx乌兰察布乌拉模组44xxx.xlsx

    文件名包含即可。

    如果没有配置 input_file_contains：

        保持原来的逻辑：
        使用第一个 xlsx 文件。

    自动忽略：

        ~$xxx.xlsx
    """

    excel_files = [
        file
        for file in input_dir.glob("*.xlsx")
        if not file.name.startswith("~$")
    ]

    # ======================================================
    # 没有 Excel
    # ======================================================

    if not excel_files:

        print()
        print(
            "[错误] input 目录没有找到 Excel 文件"
        )

        print(
            f"目录：{input_dir}"
        )

        print()

        raise Exception("处理失败")

    # ======================================================
    # 配置了文件名模糊匹配
    # ======================================================

    if input_file_contains:

        matched_files = [
            file
            for file in excel_files
            if input_file_contains.lower()
            in file.stem.lower()
        ]

        # --------------------------------------------------
        # 没找到
        # --------------------------------------------------

        if not matched_files:

            print()
            print(
                "[错误] 没有找到符合条件的 Excel 文件"
            )

            print(
                f"文件名包含：{input_file_contains}"
            )

            print(
                f"目录：{input_dir}"
            )

            print()

            print(
                "当前目录中的 Excel："
            )

            for file in excel_files:
                print(
                    f"  - {file.name}"
                )

            print()

            raise Exception("处理失败")

        # --------------------------------------------------
        # 找到多个
        # --------------------------------------------------

        if len(matched_files) > 1:

            print()
            print(
                "[错误] 找到多个符合条件的 Excel 文件："
            )

            for file in matched_files:
                print(
                    f"  - {file.name}"
                )

            print()

            print(
                "请保证 input 目录中只有一个"
                "符合 input_file_contains 的文件。"
            )

            print()

            raise Exception("处理失败")

        return matched_files[0]

    # ======================================================
    # 没配置模糊匹配
    # ======================================================

    if len(excel_files) > 1:

        print()
        print(
            "[警告] input 目录发现多个 Excel 文件："
        )

        for file in excel_files:
            print(
                f"  - {file.name}"
            )

        print()

        print(
            "当前程序默认使用第一个 Excel 文件："
        )

        print(
            f"  {excel_files[0].name}"
        )

        print()

    return excel_files[0]


# ==========================================================
# 检查 Excel 字段
# ==========================================================

def validate_columns(
    df,
    required_columns
):
    """
    检查输入 Excel 是否存在配置中需要的字段。
    """

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:

        print()
        print(
            "[错误] Excel 缺少以下字段："
        )

        for column in missing_columns:

            print(
                f"  - {column}"
            )

        print()

        print(
            "当前 Excel 的字段："
        )

        for column in df.columns:

            print(
                f"  - {column}"
            )

        print()

        return False

    return True


# ==========================================================
# 根据 group 筛选数据
# ==========================================================

def filter_group(
    df,
    group
):
    """
    根据 config.yaml 中 group 的 match 条件筛选数据。
    """

    conditions = group.get(
        "match",
        {}
    )

    # ======================================================
    # 如果没有配置 match
    # 表示这一组不过滤，全部数据都可以进入
    # ======================================================

    if not conditions:

        return df.copy()

    matched = df[
        df.apply(
            lambda row:
            row_matches(
                row,
                conditions
            ),
            axis=1
        )
    ].copy()

    return matched


# ==========================================================
# 获取需要写入模板的数据
# ==========================================================

def get_values(
    df,
    group
):
    """
    从筛选结果中获取真正需要写入目标 Excel 的字段。

    默认：

        value_column: "设备完整编号"
    """

    value_column = group.get(
        "value_column",
        "设备完整编号"
    )

    if value_column not in df.columns:

        raise ValueError(
            f"输入 Excel 不存在字段：{value_column}"
        )

    values = []

    for value in df[value_column]:

        value = clean_value(
            value
        )

        # 跳过空值
        if value:

            values.append(
                value
            )

    return values


# ==========================================================
# 写入一个 group
# ==========================================================

def write_group(
    worksheet,
    group,
    values,
    start_row
):
    """
    将一个 group 写入目标 Sheet。

    写入结构：

        description
        设备编号
        设备编号
        设备编号

    例如：

        start_row = 2

        W2  输入断路器
        W3  ID001
        W4  ID002
        W5  ID003

    blank_rows_after = 1

        下一组从 W7 开始。

    注意：

        不新增行
        不新增列
    """

    output_column = group[
        "output_column"
    ]

    # ======================================================
    # description
    #
    # description 是真正写入 Excel 的字段
    # ======================================================

    description = group.get(
        "description",
        ""
    )

    # ======================================================
    # 1. 写 description
    # ======================================================

    description_cell = (
        f"{output_column}{start_row}"
    )

    worksheet[
        description_cell
    ] = description

    # ======================================================
    # 2. 获取列号
    #
    # 例如：
    #
    # W -> 23
    # X -> 24
    # Y -> 25
    # ======================================================

    column_index = worksheet[
        f"{output_column}1"
    ].column

    # ======================================================
    # 3. 从 description 下一行开始写设备编号
    # ======================================================

    current_row = (
        start_row + 1
    )

    for value in values:

        worksheet.cell(
            row=current_row,
            column=column_index
        ).value = value

        current_row += 1

    # ======================================================
    # 4. 计算这一组结束后的位置
    # ======================================================

    blank_rows_after = int(
        group.get(
            "blank_rows_after",
            0
        )
    )

    next_start_row = (
        current_row
        + blank_rows_after
    )

    return next_start_row


# ==========================================================
# 获取 / 准备输出 Sheet
# ==========================================================

def prepare_output_sheet(
    workbook,
    output_sheet
):
    """
    准备最终输出 Sheet。

    规则：

    1. 如果目标 Sheet 已经存在：
       直接使用。

    2. 如果不存在，并且模板只有一个 Sheet：
       将唯一 Sheet 重命名为 output_sheet。

    3. 如果模板有多个 Sheet，但目标不存在：
       不自动猜测，不新增 Sheet，直接报错。

    注意：

        不自动新增 Sheet。
    """

    # ======================================================
    # 目标 Sheet 已经存在
    # ======================================================

    if output_sheet in workbook.sheetnames:

        return workbook[
            output_sheet
        ]

    # ======================================================
    # 模板只有一个 Sheet
    # 可以安全重命名
    # ======================================================

    if len(workbook.sheetnames) == 1:

        old_sheet_name = (
            workbook.sheetnames[0]
        )

        worksheet = workbook[
            old_sheet_name
        ]

        worksheet.title = (
            output_sheet
        )

        print()
        print(
            "[Sheet重命名]"
        )

        print(
            f"  {old_sheet_name}"
            f" -> "
            f"{output_sheet}"
        )

        return worksheet

    # ======================================================
    # 多个 Sheet，但目标不存在
    # ======================================================

    print()
    print(
        "[错误] 输出 Sheet 不存在："
        f"{output_sheet}"
    )

    print()

    print(
        "模板当前 Sheet："
    )

    for sheet_name in workbook.sheetnames:

        print(
            f"  - {sheet_name}"
        )

    print()

    print(
        "模板存在多个 Sheet，"
        "程序不会自动猜测需要重命名哪个 Sheet。"
    )

    print()

    print(
        "请确认模板中存在："
        f"{output_sheet}"
    )

    print()

    raise Exception("处理失败")


# ==========================================================
# 主程序
# ==========================================================

def main():

    print()
    print(
        "=" * 70
    )

    print(
        "Excel设备台账 → 定制化模组设备关系"
    )

    print(
        "=" * 70
    )

    print()

    # ======================================================
    # 1. 读取 config.yaml
    # ======================================================

    config_file = (
        BASE_DIR
        / "config.yaml"
    )

    if not config_file.exists():

        print(
            "[错误] 找不到 config.yaml"
        )

        print(
            f"文件位置：{config_file}"
        )

        raise Exception("处理失败")

    try:

        with open(
            config_file,
            "r",
            encoding="utf-8"
        ) as f:

            config = yaml.safe_load(
                f
            )

    except Exception as e:

        print(
            f"[错误] config.yaml 读取失败：{e}"
        )

        raise Exception("处理失败")

    # ======================================================
    # 2. 防止 YAML 为空
    # ======================================================

    if not config:

        print(
            "[错误] config.yaml 是空文件"
        )

        raise Exception("处理失败")

    # ======================================================
    # 3. 解析 ${project_name} 等变量
    # ======================================================

    config = expand_config_variables(
        config
    )

    # ======================================================
    # 4. 获取目录
    # ======================================================

    input_dir = (
        BASE_DIR
        / config.get(
            "input_dir",
            "input"
        )
    )

    output_dir = (
        BASE_DIR
        / config.get(
            "output_dir",
            "output"
        )
    )

    # 如果 output 目录不存在则创建
    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    # ======================================================
    # 5. 获取项目名称
    # ======================================================

    project_name = config.get(
        "project_name"
    )

    if not project_name:

        print(
            "[错误] config.yaml 没有配置 project_name"
        )

        raise Exception("处理失败")

    print(
        f"[项目名称] {project_name}"
    )

    # ======================================================
    # 6. 查找输入 Excel
    # ======================================================

    input_file_contains = config.get(
        "input_file_contains"
    )

    input_file = get_input_excel(
        input_dir,
        input_file_contains
    )

    print(
        f"[输入文件] {input_file.name}"
    )

    # ======================================================
    # 7. 读取输入 Excel
    # ======================================================

    input_sheet = config.get(
        "input_sheet",
        0
    )

    output_sheet_default = config.get(
        "output_sheet"
    )

    if not output_sheet_default:

        print(
            "[错误] config.yaml 没有配置 output_sheet"
        )

        raise Exception("处理失败")

    print(
        f"[输入Sheet] {input_sheet}"
    )

    print(
        f"[输出Sheet] {output_sheet_default}"
    )

    try:

        df = pd.read_excel(
            input_file,
            sheet_name=input_sheet,
            dtype=str
        )

    except Exception as e:

        print()
        print(
            f"[错误] 输入 Excel 读取失败：{e}"
        )

        raise Exception("处理失败")

    # ======================================================
    # 8. 填充空值
    # ======================================================

    df = df.fillna("")

    # ======================================================
    # 9. 清理所有字段前后空格
    # ======================================================

    for column in df.columns:

        df[column] = df[column].map(
            clean_value
        )

    original_count = len(df)

    print(
        f"[原始数据] {original_count} 行"
    )

    # ======================================================
    # 10. 删除完全空行
    # ======================================================

    df = df.loc[
        df.astype(str).apply(
            lambda row:
            any(
                str(value).strip()
                for value in row
            ),
            axis=1
        )
    ].copy()

    print(
        f"[删除空行后] {len(df)} 行"
    )

    # ======================================================
    # 11. 全局 require_non_empty
    # ======================================================

    require_non_empty = config.get(
        "require_non_empty",
        []
    )

    for column in require_non_empty:

        if column not in df.columns:

            print()

            print(
                f"[错误] 必须字段不存在：{column}"
            )

            raise Exception("处理失败")

        df = df[
            df[column]
            .astype(str)
            .str.strip()
            != ""
        ]

    print(
        f"[全局非空过滤后] {len(df)} 行"
    )

    # ======================================================
    # 12. 获取 groups
    # ======================================================

    groups = config.get(
        "groups",
        []
    )

    if not groups:

        print()

        print(
            "[错误] config.yaml 没有配置 groups"
        )

        raise Exception("处理失败")

    # ======================================================
    # 13. 检查所有需要的字段
    # ======================================================

    required_columns = set()

    for group in groups:

        # ----------------------------------------------
        # match 中的字段
        # ----------------------------------------------

        conditions = group.get(
            "match",
            {}
        )

        for column in conditions:

            required_columns.add(
                column
            )

        # ----------------------------------------------
        # value_column
        # ----------------------------------------------

        value_column = group.get(
            "value_column",
            "设备完整编号"
        )

        required_columns.add(
            value_column
        )

    if not validate_columns(
        df,
        required_columns
    ):

        raise Exception("处理失败")

    # ======================================================
    # 14. 找到输出模板
    # ======================================================

    template_file = (
        output_dir
        / config.get(
            "template_file",
            "定制化模组设备关系.xlsx"
        )
    )

    if not template_file.exists():

        print()

        print(
            "[错误] 找不到输出模板："
        )

        print(
            template_file
        )

        print()

        print(
            "请确认模板文件放在："
        )

        print(
            output_dir
        )

        raise Exception("处理失败")

    print(
        f"[输出模板] {template_file.name}"
    )

    # ======================================================
    # 15. 打开输出模板
    # ======================================================

    try:

        workbook = load_workbook(
            template_file
        )

    except Exception as e:

        print()

        print(
            "[错误] 输出模板打开失败："
        )

        print(e)

        raise Exception("处理失败")

    # ======================================================
    # 16. 准备默认输出 Sheet
    #
    # 如果模板只有一个 Sheet：
    # 自动重命名为 output_sheet。
    #
    # 不新增 Sheet。
    # ======================================================

    default_worksheet = prepare_output_sheet(
        workbook,
        output_sheet_default
    )

    # ======================================================
    # 17. 记录每个 Sheet + Column 下一次写入的位置
    #
    # 例如：
    #
    # ("乌兰察布乌拉模组44", "F") -> 15
    #
    # ("乌兰察布乌拉模组44", "W") -> 10
    #
    # F 和 W 各自独立。
    # ======================================================

    next_position = {}

    # ======================================================
    # 统计
    # ======================================================

    total_write_count = 0

    # ======================================================
    # 18. 开始处理每一个 group
    # ======================================================

    for index, group in enumerate(
        groups,
        start=1
    ):

        group_name = group.get(
            "name",
            f"筛选组{index}"
        )

        # ==================================================
        # output_sheet
        #
        # group 可以自己覆盖。
        #
        # 如果没有配置：
        # 使用全局 output_sheet。
        # ==================================================

        output_sheet = group.get(
            "output_sheet",
            output_sheet_default
        )

        output_column = group.get(
            "output_column"
        )

        # ==================================================
        # 检查基本配置
        # ==================================================

        if not output_sheet:

            print()

            print(
                f"[错误] 分组「{group_name}」"
                "没有配置 output_sheet"
            )

            continue

        if not output_column:

            print()

            print(
                f"[错误] 分组「{group_name}」"
                "没有配置 output_column"
            )

            continue

        # ==================================================
        # 当前 Sheet + Column
        #
        # 这是判断“上一组在哪里结束”的关键
        # ==================================================

        position_key = (
            output_sheet,
            output_column
        )

        # ==================================================
        # 19. 判断这一组从哪里开始
        # ==================================================

        if group.get(
            "auto_position",
            False
        ):

            # ------------------------------------------------
            # auto_position = true
            #
            # 从同 Sheet + 同列的上一组结束位置开始
            # ------------------------------------------------

            if position_key not in next_position:

                print()

                print(
                    f"[错误] 分组「{group_name}」"
                )

                print(
                    "配置了 auto_position: true"
                )

                print(
                    "但是没有找到前一个相同"
                    "Sheet + Column 的分组。"
                )

                print(
                    f"Sheet：{output_sheet}"
                )

                print(
                    f"列：{output_column}"
                )

                continue

            start_row = next_position[
                position_key
            ]

        else:

            # ------------------------------------------------
            # auto_position 没有开启
            #
            # 使用自己的 start_row
            # ------------------------------------------------

            start_row = int(
                group.get(
                    "start_row",
                    2
                )
            )

        # ==================================================
        # 20. 打印分组信息
        # ==================================================

        print()

        print(
            "-" * 70
        )

        print(
            f"[分组 {index}] {group_name}"
        )

        print(
            f"目标Sheet：{output_sheet}"
        )

        print(
            f"目标列：{output_column}"
        )

        print(
            f"开始行：{start_row}"
        )

        print(
            f"描述：{group.get('description', '')}"
        )

        # ==================================================
        # 21. 检查目标 Sheet
        # ==================================================

        if output_sheet not in workbook.sheetnames:

            print(
                f"[错误] 输出模板不存在 Sheet："
                f"{output_sheet}"
            )

            continue

        worksheet = workbook[
            output_sheet
        ]

        # ==================================================
        # 22. 执行筛选
        # ==================================================

        matched = filter_group(
            df,
            group
        )

        print(
            f"筛选匹配：{len(matched)} 行"
        )

        # ==================================================
        # 23. 获取要写入的数据
        # ==================================================

        try:

            values = get_values(
                matched,
                group
            )

        except Exception as e:

            print(
                f"[错误] 获取写入数据失败：{e}"
            )

            continue

        print(
            f"实际写入：{len(values)} 个"
        )

        # ==================================================
        # 24. 写入模板
        #
        # 写入顺序：
        #
        # description
        # 设备编号1
        # 设备编号2
        # 设备编号3
        #
        # 不新增行。
        # 不新增列。
        # ==================================================

        next_row = write_group(
            worksheet,
            group,
            values,
            start_row
        )

        # ==================================================
        # 25. 记录下一组位置
        #
        # 只有同 Sheet + 同 Column
        # 才会使用这个位置。
        # ==================================================

        next_position[
            position_key
        ] = next_row

        total_write_count += len(
            values
        )

        print(
            "本组写入完成"
        )

        print(
            f"下一组自动从："
            f"{output_column}{next_row}"
            "开始"
        )

    # ======================================================
    # 26. 另存为输出文件
    #
    # 默认：
    #
    # output/{project_name}.xlsx
    #
    # 也可以 config 中配置：
    #
    # output_file: "xxx.xlsx"
    #
    # 不覆盖 template_file。
    # ======================================================

    output_file = config.get(
        "output_file",
        f"{project_name}.xlsx"
    )

    output_path = (
        output_dir
        / output_file
    )

    # ======================================================
    # 防止 output_file 与模板相同
    # ======================================================

    if (
        output_path.resolve()
        ==
        template_file.resolve()
    ):

        print()

        print(
            "[错误] output_file 不能与 "
            "template_file 相同"
        )

        print(
            f"模板：{template_file}"
        )

        print(
            f"输出：{output_path}"
        )

        print()

        raise Exception("处理失败")

    # ======================================================
    # 保存
    # ======================================================

    try:

        workbook.save(
            output_path
        )

    except Exception as e:

        print()

        print(
            "[错误] 输出 Excel 保存失败："
        )

        print(e)

        raise Exception("处理失败")

    # ======================================================
    # 27. 最终统计
    # ======================================================

    print()

    print(
        "=" * 70
    )

    print(
        "处理完成"
    )

    print(
        "=" * 70
    )

    print(
        f"输入文件：{input_file.name}"
    )

    print(
        f"输出文件：{output_path.name}"
    )

    print(
        f"原始数据：{original_count} 行"
    )

    print(
        f"最终有效数据：{len(df)} 行"
    )

    print(
        f"总写入数量：{total_write_count} 个"
    )

    print()

    print(
        f"文件位置：{output_path}"
    )

    print()


# ==========================================================
# 程序入口
# ==========================================================

if __name__ == "__main__":
    main()