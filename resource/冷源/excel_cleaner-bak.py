from pathlib import Path
import sys

import pandas as pd
import yaml
from openpyxl import load_workbook


# ==========================================================
# 基础目录
# ==========================================================

BASE_DIR = Path(__file__).resolve().parent


# ==========================================================
# 工具函数
# ==========================================================

def clean_value(value):
    """
    统一处理Excel单元格值

    None / NaN -> ""
    其他 -> 字符串并去掉首尾空格
    """

    if pd.isna(value):
        return ""

    return str(value).strip()


# ==========================================================
# 判断一行是否符合筛选条件
# ==========================================================

def row_matches(row, conditions):
    """
    判断一行数据是否满足 match 条件。

    默认所有条件都是 AND。

    例如：

    match:
      应用类型: "输入断路器"
      所属整机类型: "一体冷源"

    表示：

    应用类型 = 输入断路器
    AND
    所属整机类型 = 一体冷源
    """

    for column, expected in conditions.items():

        actual = clean_value(
            row.get(column, "")
        )

        # ==================================================
        # 如果配置的是列表
        #
        # 例如：
        #
        # 应用类型:
        #   - "输入断路器"
        #   - "输出断路器"
        #
        # 表示 OR
        # ==================================================

        if isinstance(expected, list):

            expected_values = [
                clean_value(value)
                for value in expected
            ]

            if actual not in expected_values:
                return False

        else:

            expected_value = clean_value(
                expected
            )

            if actual != expected_value:
                return False

    return True


# ==========================================================
# 查找输入 Excel
# ==========================================================

def get_input_excel(input_dir):
    """
    从 input 目录查找 Excel。

    默认取第一个 .xlsx 文件。

    会自动忽略 Excel 临时文件：
    ~$xxx.xlsx
    """

    excel_files = [
        file
        for file in input_dir.glob("*.xlsx")
        if not file.name.startswith("~$")
    ]

    if not excel_files:

        print()
        print("[错误] input 目录没有找到 Excel 文件")
        print(f"目录：{input_dir}")
        print()

        sys.exit(1)

    # 如果有多个Excel，提示用户
    if len(excel_files) > 1:

        print()
        print("[警告] input 目录发现多个 Excel 文件：")

        for file in excel_files:
            print(f"  - {file.name}")

        print()
        print("当前程序默认使用第一个 Excel 文件：")
        print(f"  {excel_files[0].name}")
        print()

    return excel_files[0]


# ==========================================================
# 检查 Excel 字段
# ==========================================================

def validate_columns(df, required_columns):
    """
    检查输入 Excel 是否存在配置中需要的字段。
    """

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:

        print()
        print("[错误] Excel 缺少以下字段：")

        for column in missing_columns:
            print(f"  - {column}")

        print()
        print("当前 Excel 的字段：")

        for column in df.columns:
            print(f"  - {column}")

        print()

        return False

    return True


# ==========================================================
# 根据 group 筛选数据
# ==========================================================

def filter_group(df, group):
    """
    根据 config.yaml 中 group 的 match 条件筛选数据。
    """

    conditions = group.get(
        "match",
        {}
    )

    # 如果没有配置 match
    # 表示这一组不过滤，全部数据都可以进入
    if not conditions:

        return df.copy()

    matched = df[
        df.apply(
            lambda row:
            row_matches(
                row,
                conditions
            ),
            axis=1
        )
    ].copy()

    return matched


# ==========================================================
# 获取需要写入模板的数据
# ==========================================================

def get_values(df, group):
    """
    从筛选结果中获取真正需要写入目标 Excel 的字段。

    config:

    value_column: "设备完整编号"
    """

    value_column = group.get(
        "value_column",
        "设备完整编号"
    )

    if value_column not in df.columns:

        raise ValueError(
            f"输入 Excel 不存在字段：{value_column}"
        )

    values = []

    for value in df[value_column]:

        value = clean_value(value)

        # 跳过空值
        if value:
            values.append(value)

    return values


# ==========================================================
# 写入一个 group
# ==========================================================

def write_group(
    worksheet,
    group,
    values,
    start_row
):
    """
    将一个 group 写入目标 Sheet。

    例如：

    start_row = 2

    W2  输入断路器（一体冷源01/02逆变）
    W3  ID001
    W4  ID002
    W5  ID003
    W6  ID004

    blank_rows_after = 1

    那么下一组从 W8 开始。
    """

    output_column = group[
        "output_column"
    ]

    description = group.get(
        "description",
        ""
    )

    # ======================================================
    # 1. 写 description
    # ======================================================

    description_cell = (
        f"{output_column}{start_row}"
    )

    worksheet[
        description_cell
    ] = description

    # ======================================================
    # 2. 获取列号
    #
    # W -> 23
    # X -> 24
    # Y -> 25
    # ======================================================

    column_index = worksheet[
        f"{output_column}1"
    ].column

    # ======================================================
    # 3. 从 description 下一行开始写设备编号
    # ======================================================

    current_row = start_row + 1

    for value in values:

        worksheet.cell(
            row=current_row,
            column=column_index
        ).value = value

        current_row += 1

    # ======================================================
    # 4. 计算这一组结束后的位置
    # ======================================================

    blank_rows_after = int(
        group.get(
            "blank_rows_after",
            0
        )
    )

    next_start_row = (
        current_row
        + blank_rows_after
    )

    return next_start_row


# ==========================================================
# 主程序
# ==========================================================

def main():

    print()
    print("=" * 70)
    print("Excel设备台账 → 定制化模组设备关系")
    print("=" * 70)
    print()

    # ======================================================
    # 1. 读取 config.yaml
    # ======================================================

    config_file = (
        BASE_DIR
        / "config.yaml"
    )

    if not config_file.exists():

        print("[错误] 找不到 config.yaml")

        sys.exit(1)

    try:

        with open(
            config_file,
            "r",
            encoding="utf-8"
        ) as f:

            config = yaml.safe_load(f)

    except Exception as e:

        print(
            f"[错误] config.yaml 读取失败：{e}"
        )

        sys.exit(1)

    # 防止yaml为空
    if not config:

        print("[错误] config.yaml 是空文件")

        sys.exit(1)

    # ======================================================
    # 2. 获取目录
    # ======================================================

    input_dir = (
        BASE_DIR
        / config.get(
            "input_dir",
            "input"
        )
    )

    output_dir = (
        BASE_DIR
        / config.get(
            "output_dir",
            "output"
        )
    )

    # 如果output目录不存在则创建
    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    # ======================================================
    # 3. 查找输入Excel
    # ======================================================

    input_file = get_input_excel(
        input_dir
    )

    print(
        f"[输入文件] {input_file.name}"
    )

    # ======================================================
    # 4. 读取输入Excel
    # ======================================================

    input_sheet = config.get(
        "input_sheet",
        0
    )

    try:

        df = pd.read_excel(
            input_file,
            sheet_name=input_sheet,
            dtype=str
        )

    except Exception as e:

        print()
        print(
            f"[错误] 输入 Excel 读取失败：{e}"
        )

        sys.exit(1)

    # ======================================================
    # 5. 填充空值
    # ======================================================

    df = df.fillna("")

    # ======================================================
    # 6. 清理所有字段前后空格
    # ======================================================

    for column in df.columns:

        df[column] = df[column].map(
            clean_value
        )

    original_count = len(df)

    print(
        f"[原始数据] {original_count} 行"
    )

    # ======================================================
    # 7. 删除完全空行
    # ======================================================

    df = df.loc[
        df.astype(str).apply(
            lambda row:
            any(
                str(value).strip()
                for value in row
            ),
            axis=1
        )
    ].copy()

    print(
        f"[删除空行后] {len(df)} 行"
    )

    # ======================================================
    # 8. 全局 require_non_empty
    # ======================================================

    require_non_empty = config.get(
        "require_non_empty",
        []
    )

    for column in require_non_empty:

        if column not in df.columns:

            print()
            print(
                f"[错误] 必须字段不存在：{column}"
            )

            sys.exit(1)

        df = df[
            df[column]
            .astype(str)
            .str.strip()
            != ""
        ]

    print(
        f"[全局非空过滤后] {len(df)} 行"
    )

    # ======================================================
    # 9. 获取 groups
    # ======================================================

    groups = config.get(
        "groups",
        []
    )

    if not groups:

        print()
        print(
            "[错误] config.yaml 没有配置 groups"
        )

        sys.exit(1)

    # ======================================================
    # 10. 检查所有需要的字段
    # ======================================================

    required_columns = set()

    for group in groups:

        # ----------------------------------------------
        # match 中的字段
        # ----------------------------------------------

        conditions = group.get(
            "match",
            {}
        )

        for column in conditions:

            required_columns.add(
                column
            )

        # ----------------------------------------------
        # value_column
        # ----------------------------------------------

        value_column = group.get(
            "value_column",
            "设备完整编号"
        )

        required_columns.add(
            value_column
        )

    if not validate_columns(
        df,
        required_columns
    ):

        sys.exit(1)

    # ======================================================
    # 11. 找到输出模板
    # ======================================================

    template_file = (
        output_dir
        / config.get(
            "template_file",
            "定制化模组设备关系.xlsx"
        )
    )

    if not template_file.exists():

        print()
        print(
            "[错误] 找不到输出模板："
        )

        print(
            template_file
        )

        print()
        print(
            "请确认模板文件放在："
        )

        print(
            output_dir
        )

        sys.exit(1)

    print(
        f"[输出模板] {template_file.name}"
    )

    # ======================================================
    # 12. 打开输出模板
    # ======================================================

    try:

        workbook = load_workbook(
            template_file
        )

    except Exception as e:

        print()
        print(
            "[错误] 输出模板打开失败："
        )

        print(e)

        sys.exit(1)

    # ======================================================
    # 13. 记录每个 Sheet + Column 下一次写入的位置
    #
    # 非常重要
    #
    # 例如：
    #
    # ("韶关百旺32液冷", "W") -> 15
    #
    # 表示下一组W列从15行开始。
    #
    # ("韶关百旺32液冷", "X") -> 10
    #
    # 表示下一组X列从10行开始。
    # ======================================================

    next_position = {}

    # ======================================================
    # 统计
    # ======================================================

    total_write_count = 0

    # ======================================================
    # 14. 开始处理每一个 group
    # ======================================================

    for index, group in enumerate(
        groups,
        start=1
    ):

        group_name = group.get(
            "name",
            f"筛选组{index}"
        )

        output_sheet = group.get(
            "output_sheet"
        )

        output_column = group.get(
            "output_column"
        )

        # ==================================================
        # 检查基本配置
        # ==================================================

        if not output_sheet:

            print()
            print(
                f"[错误] 分组「{group_name}」"
                f"没有配置 output_sheet"
            )

            continue

        if not output_column:

            print()
            print(
                f"[错误] 分组「{group_name}」"
                f"没有配置 output_column"
            )

            continue

        # ==================================================
        # 当前 Sheet + Column
        #
        # 这是判断“上一组在哪里结束”的关键
        # ==================================================

        position_key = (
            output_sheet,
            output_column
        )

        # ==================================================
        # 15. 判断这一组从哪里开始
        # ==================================================

        if group.get(
            "auto_position",
            False
        ):

            # ------------------------------------------------
            # auto_position = true
            #
            # 从同Sheet + 同列的上一组结束位置开始
            # ------------------------------------------------

            if position_key not in next_position:

                print()
                print(
                    f"[错误] 分组「{group_name}」"
                )

                print(
                    "配置了 auto_position: true"
                )

                print(
                    "但是没有找到前一个相同"
                    "Sheet + Column 的分组。"
                )

                print(
                    f"Sheet：{output_sheet}"
                )

                print(
                    f"列：{output_column}"
                )

                continue

            start_row = next_position[
                position_key
            ]

        else:

            # ------------------------------------------------
            # auto_position 没有开启
            #
            # 使用自己配置的 start_row
            # ------------------------------------------------

            start_row = int(
                group.get(
                    "start_row",
                    2
                )
            )

        # ==================================================
        # 16. 打印分组信息
        # ==================================================

        print()
        print(
            "-" * 70
        )

        print(
            f"[分组 {index}] {group_name}"
        )

        print(
            f"目标Sheet：{output_sheet}"
        )

        print(
            f"目标列：{output_column}"
        )

        print(
            f"开始行：{start_row}"
        )

        print(
            f"描述："
            f"{group.get('description', '')}"
        )

        # ==================================================
        # 17. 检查目标 Sheet
        # ==================================================

        if output_sheet not in workbook.sheetnames:

            print(
                f"[错误] 输出模板不存在 Sheet："
                f"{output_sheet}"
            )

            continue

        worksheet = workbook[
            output_sheet
        ]

        # ==================================================
        # 18. 执行筛选
        # ==================================================

        matched = filter_group(
            df,
            group
        )

        print(
            f"筛选匹配：{len(matched)} 行"
        )

        # ==================================================
        # 19. 获取要写入的数据
        # ==================================================

        try:

            values = get_values(
                matched,
                group
            )

        except Exception as e:

            print(
                f"[错误] 获取写入数据失败：{e}"
            )

            continue

        print(
            f"实际写入：{len(values)} 个"
        )

        # ==================================================
        # 20. 写入模板
        # ==================================================

        next_row = write_group(
            worksheet,
            group,
            values,
            start_row
        )

        # ==================================================
        # 21. 记录下一组位置
        #
        # 这是防止第二组覆盖第一组的核心
        # ==================================================

        next_position[
            position_key
        ] = next_row

        total_write_count += len(values)

        print(
            f"本组写入完成"
        )

        print(
            f"下一组自动从："
            f"{output_column}{next_row}"
            f"开始"
        )

    # ======================================================
    # 22. 保存模板
    # ======================================================

    try:

        workbook.save(
            template_file
        )

    except Exception as e:

        print()
        print(
            "[错误] 输出Excel保存失败："
        )

        print(e)

        sys.exit(1)

    # ======================================================
    # 23. 最终统计
    # ======================================================

    print()
    print("=" * 70)
    print("处理完成")
    print("=" * 70)

    print(
        f"输入文件：{input_file.name}"
    )

    print(
        f"输出文件：{template_file.name}"
    )

    print(
        f"原始数据：{original_count} 行"
    )

    print(
        f"最终有效数据：{len(df)} 行"
    )

    print(
        f"总写入数量：{total_write_count} 个"
    )

    print()
    print(
        f"文件位置：{template_file}"
    )

    print()


# ==========================================================
# 程序入口
# ==========================================================

if __name__ == "__main__":
    main()